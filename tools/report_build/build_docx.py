# -*- coding: utf-8 -*-
"""
Render the working-draft report to DOCX using python-docx.

The document is built from real, reusable Word styles rather than from directly
formatted paragraphs, so the output stays editable in Word.

Usage
-----
    python3 build_docx.py OUT.docx [page_map.json]

If a page map is supplied, the contents, list of figures and list of tables are
written with page numbers. build.sh produces that map from a first-pass PDF and
then re-runs this script, which is how the front matter ends up correct without
relying on a field update at open time.
"""
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import report_content as C

# A4 minus 2.5 cm margins on each side.
TEXT_COL_IN = (21.0 - 5.0) / 2.54

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))

# Times New Roman is present on the build machine, so the PDF matches the DOCX
# rather than relying on a metric substitute. No font file is embedded or
# redistributed by this script.
SERIF = "Times New Roman"
BODY_PT = 11
GREY = RGBColor(0x44, 0x44, 0x44)

# Set True to render the contents as static text instead of a Word field.
STATIC_TOC = True


# ---------------------------------------------------------------- low level

def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(k), v)
    return e


def set_cell_border(cell, sz=4, color="999999", edges=("top", "bottom", "left", "right")):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in edges:
        e = borders.find(qn("w:" + edge))
        if e is None:
            e = OxmlElement("w:" + edge)
            borders.append(e)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:color"), color)


def clear_cell_borders(cell, keep=()):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom", "left", "right"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single" if edge in keep else "nil")
        if edge in keep:
            e.set(qn("w:sz"), "6")
            e.set(qn("w:color"), "666666")
        borders.append(e)
    tcPr.append(borders)


def shade_cell(cell, fill="F2F2F2"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(_el("w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill}))


def keep_with_next(par, on=True):
    par.paragraph_format.keep_with_next = on


def add_field(par, instr):
    """Insert a Word field such as PAGE or TOC."""
    r1 = par.add_run()
    r1._r.append(_el("w:fldChar", **{"w:fldCharType": "begin"}))
    r2 = par.add_run()
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    r2._r.append(t)
    r3 = par.add_run()
    r3._r.append(_el("w:fldChar", **{"w:fldCharType": "separate"}))
    r4 = par.add_run("1")
    r5 = par.add_run()
    r5._r.append(_el("w:fldChar", **{"w:fldCharType": "end"}))


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:tblHeader", **{"w:val": "true"}))


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(_el("w:cantSplit"))



def force_font(style, name=SERIF):
    """Pin a style to a real font, clearing any theme reference.

    Word's built-in heading styles point at the theme's major font rather than
    at a named typeface. Setting style.font.name alone leaves that reference in
    place, so headings come out in the theme sans-serif. Removing the *Theme
    attributes and writing the name explicitly fixes both Word and LibreOffice.
    """
    style.font.name = name
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        key = qn("w:" + attr)
        if key in rFonts.attrib:
            del rFonts.attrib[key]
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn("w:" + attr), name)


# ------------------------------------------------------------------ styles

def build_styles(doc):
    st = doc.styles

    normal = st["Normal"]
    force_font(normal)
    normal.font.size = Pt(BODY_PT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.20
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def heading(name, size, bold=True, before=14, after=6, italic=False):
        s = st[name]
        force_font(s)
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.italic = italic
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.line_spacing = 1.1
        return s

    heading("Heading 1", 14.5, before=18, after=7)
    heading("Heading 2", 12.0, before=13, after=5)
    heading("Heading 3", 11.0, before=11, after=4, italic=True)

    ttl = st["Title"]
    force_font(ttl)
    ttl.font.size = Pt(21)
    ttl.font.bold = True
    ttl.font.color.rgb = RGBColor(0, 0, 0)
    ttl.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ttl.paragraph_format.space_after = Pt(10)
    ttl.paragraph_format.line_spacing = 1.1

    cap = st["Caption"]
    force_font(cap)
    cap.font.size = Pt(9)
    cap.font.italic = False
    cap.font.bold = False
    cap.font.color.rgb = GREY
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cap.paragraph_format.space_before = Pt(5)
    cap.paragraph_format.space_after = Pt(11)
    cap.paragraph_format.line_spacing = 1.05

    def custom(name, base="Normal", **kw):
        try:
            s = st.add_style(name, 1)
        except Exception:
            s = st[name]
        s.base_style = st[base]
        force_font(s)
        for k, v in kw.items():
            setattr(s.font, k, v)
        return s

    s = custom("Table Caption", size=Pt(9))
    s.font.color.rgb = GREY
    s.paragraph_format.space_before = Pt(10)
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    s.paragraph_format.line_spacing = 1.05
    s.paragraph_format.keep_with_next = True

    s = custom("Table Body", size=Pt(8.7))
    s.paragraph_format.space_after = Pt(1.5)
    s.paragraph_format.space_before = Pt(1.5)
    s.paragraph_format.line_spacing = 1.03
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    s = custom("Equation Text", size=Pt(11))
    s.paragraph_format.space_before = Pt(8)
    s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    s = custom("Reference Entry", size=Pt(10))
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.left_indent = Cm(1.0)
    s.paragraph_format.first_line_indent = Cm(-1.0)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.line_spacing = 1.1

    s = custom("Front Matter Entry", size=Pt(10.5))
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    s.paragraph_format.line_spacing = 1.1

    s = custom("Placeholder Label", size=Pt(10))
    s.font.bold = True
    s.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(4)
    s.paragraph_format.space_before = Pt(4)

    s = custom("Placeholder Body", size=Pt(9))
    s.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.paragraph_format.space_after = Pt(3)
    s.paragraph_format.line_spacing = 1.1

    s = custom("Small Note", size=Pt(9.5))
    s.font.color.rgb = GREY
    s.paragraph_format.left_indent = Cm(0.7)
    s.paragraph_format.space_after = Pt(8)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for nm in ("List Bullet", "List Number"):
        s = st[nm]
        force_font(s)
        s.font.size = Pt(BODY_PT)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.line_spacing = 1.15
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def page_setup(doc):
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        for a in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(sec, a, Cm(2.5))
        sec.header_distance = Cm(1.4)
        sec.footer_distance = Cm(1.4)


def decorate_section(sec, header_text, numbering=True):
    hp = sec.header.paragraphs[0]
    hp.text = ""
    if header_text:
        r = hp.add_run(header_text)
        r.font.size = Pt(8.5)
        r.font.name = SERIF
        r.font.color.rgb = GREY
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if numbering:
        add_field(fp, " PAGE ")
        for r in fp.runs:
            r.font.size = Pt(9)
            r.font.name = SERIF
            r.font.color.rgb = GREY


# --------------------------------------------------------------- rendering

class Builder:
    def __init__(self, doc, page_map=None):
        self.doc = doc
        self.page_map = page_map or {}
        self.fig_n = 0
        self.tab_n = 0
        self.figures = []   # (label, short caption, key)
        self.tables = []
        self.h1 = 0
        self.h2 = 0
        self.toc = []       # (level, text)
        self.appendix = None
        self.afig = {}
        self.atab = {}

    # -- helpers
    def para(self, text, style=None):
        p = self.doc.add_paragraph(text, style=style)
        return p

    def pnum(self, key):
        v = self.page_map.get(key)
        return str(v) if v else ""

    # -- blocks
    def heading1(self, text):
        self.h1 += 1
        self.h2 = 0
        label = "%d" % self.h1
        full = "%s  %s" % (label, text)
        p = self.doc.add_heading(full, level=1)
        self.toc.append((1, full, "h1:%d" % self.h1))
        return p

    def heading2(self, text):
        self.h2 += 1
        label = "%d.%d" % (self.h1, self.h2)
        full = "%s  %s" % (label, text)
        p = self.doc.add_heading(full, level=2)
        self.toc.append((2, full, "h2:%s" % label))
        return p

    def appendix_heading(self, letter, text):
        self.appendix = letter
        full = "Appendix %s  %s" % (letter, text)
        p = self.doc.add_heading(full, level=1)
        self.toc.append((1, full, "app:%s" % letter))
        return p

    def equation(self, text, number):
        tbl = self.doc.add_table(rows=1, cols=2)
        tbl.autofit = False
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        widths = [5.4, 0.7]
        for i, w in enumerate(widths):
            tbl.columns[i].width = Inches(w)
        c0, c1 = tbl.rows[0].cells
        for c in (c0, c1):
            clear_cell_borders(c)
        c0.width, c1.width = Inches(widths[0]), Inches(widths[1])
        p = c0.paragraphs[0]
        p.style = self.doc.styles["Equation Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # text is a list of (glyphs, role) tokens: "i" italic variable,
        # "r" upright, "sub" italic subscript. Keeping variables italic and
        # operators upright is the usual convention and reads correctly in both
        # Word and the PDF without an equation editor.
        tokens = text if isinstance(text, list) else [(text, "i")]
        for glyphs, role in tokens:
            r = p.add_run(glyphs)
            r.font.size = Pt(11.5)
            if role == "i":
                r.font.italic = True
            elif role == "sub":
                r.font.italic = True
                r.font.subscript = True
        p2 = c1.paragraphs[0]
        p2.style = self.doc.styles["Equation Text"]
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p2.add_run("(%s)" % number).font.size = Pt(10.5)

    def figure(self, spec, label=None, key=None):
        self.fig_n += 1 if label is None else 0
        lab = label or ("Figure %d" % self.fig_n)
        holder = self.doc.add_paragraph()
        holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        holder.paragraph_format.space_before = Pt(9)
        holder.paragraph_format.space_after = Pt(0)
        keep_with_next(holder)

        path = spec.get("path")
        if path:
            full = os.path.join(REPO, path)
            holder.add_run().add_picture(full, height=Inches(spec["height"]))
        else:
            self._placeholder(holder, lab, spec)

        cap = self.doc.add_paragraph(style="Caption")
        cap.add_run("%s. " % lab).bold = True
        cap.add_run(spec["caption"])
        short = spec["caption"].split(".")[0].strip()
        self.figures.append((lab, short, key or lab))
        return lab

    def _placeholder(self, holder, lab, spec):
        # A bordered box of realistic figure dimensions, so that page layout
        # stays honest before the image exists.
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        w = 5.9
        tbl.columns[0].width = Inches(w)
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(w)
        set_cell_border(cell, sz=8, color="AAAAAA")
        shade_cell(cell, "FAFAFA")
        no_split(tbl.rows[0])
        tr = tbl.rows[0]
        tr.height = Inches(spec["height"])

        p0 = cell.paragraphs[0]
        p0.style = self.doc.styles["Placeholder Label"]
        p0.add_run(lab.upper())
        p1 = cell.add_paragraph(style="Placeholder Body")
        p1.add_run("Figure required. Not yet available.").italic = True
        p2 = cell.add_paragraph(style="Placeholder Body")
        p2.add_run(spec["need"])
        # move the table into the holder position
        holder._p.addnext(tbl._tbl)
        holder.add_run("")

    def table(self, spec, label=None, key=None):
        if label is None:
            self.tab_n += 1
            lab = "Table %d" % self.tab_n
        else:
            lab = label
        cap = self.doc.add_paragraph(style="Table Caption")
        cap.add_run("%s. " % lab).bold = True
        cap.add_run(spec["caption"])

        headers = spec["headers"]
        rows = spec["rows"]
        widths = spec.get("widths")
        tbl = self.doc.add_table(rows=1, cols=len(headers))
        tbl.style = "Table Grid"
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        hdr = tbl.rows[0]
        repeat_header(hdr)
        for i, h in enumerate(headers):
            c = hdr.cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.style = self.doc.styles["Table Body"]
            p.add_run(h).bold = True
            shade_cell(c, "EDEDED")
        for r in rows:
            tr = tbl.add_row()
            no_split(tr)
            for i, v in enumerate(r):
                c = tr.cells[i]
                c.text = ""
                p = c.paragraphs[0]
                p.style = self.doc.styles["Table Body"]
                p.add_run(str(v))
        if widths:
            # Scale the requested widths to the text column so that a table can
            # never bleed into the margins, whatever the content module asks for.
            total = float(sum(widths))
            if total > TEXT_COL_IN:
                widths = [w * TEXT_COL_IN / total for w in widths]
            for i, w in enumerate(widths):
                for row in tbl.rows:
                    row.cells[i].width = Inches(w)
                tbl.columns[i].width = Inches(w)
        for row in tbl.rows:
            for c in row.cells:
                clear_cell_borders(c, keep=("top", "bottom"))
        # A short table that straddles a page boundary reads badly, and header
        # repetition does not help when only one or two rows carry over. Binding
        # every row but the last to the next one pushes the whole table onto the
        # following page instead. Long tables are left to break naturally, since
        # they cannot fit on one page and their headers do repeat.
        if len(rows) <= 12:
            for row in tbl.rows[:-1]:
                for c in row.cells:
                    for par in c.paragraphs:
                        par.paragraph_format.keep_with_next = True
        after = self.doc.add_paragraph()
        after.paragraph_format.space_after = Pt(8)
        short = spec["caption"].split(".")[0].strip()
        self.tables.append((lab, short, key or lab))
        return lab

    def render(self, blocks, fig_src, tab_src, fig_label=None, tab_label=None):
        for b in blocks:
            kind = b[0]
            if kind == "h1":
                self.heading1(b[1])
            elif kind == "h2":
                self.heading2(b[1])
            elif kind == "p":
                self.para(b[1])
            elif kind == "note":
                self.para(b[1], style="Small Note")
            elif kind == "bul":
                for it in b[1]:
                    self.para(it, style="List Bullet")
            elif kind == "num":
                for it in b[1]:
                    self.para(it, style="List Number")
            elif kind == "eq":
                self.equation(b[1], b[2])
            elif kind == "fig":
                spec = fig_src[b[1]] if b[1] in fig_src else C.APPENDIX_FIGURES[b[1]]
                lab = fig_label(b[1]) if fig_label else None
                self.figure(spec, label=lab, key=b[1])
            elif kind == "afig":
                spec = C.APPENDIX_FIGURES[b[1]]
                lab = fig_label(b[1]) if fig_label else None
                self.figure(spec, label=lab, key=b[1])
            elif kind == "tab":
                self.table(tab_src[b[1]], key=b[1])
            elif kind == "atab":
                spec = C.APPENDIX_TABLES[b[1]]
                lab = tab_label(b[1]) if tab_label else None
                self.table(spec, label=lab, key=b[1])
            elif kind == "pb":
                self.doc.add_page_break()


# ------------------------------------------------------------- front matter

APP_FIG_LABEL = {
    "convergence_appendix": "Figure C.1",
    "biasT_2020": "Figure D.1",
    "magnitude_appendix": "Figure E.1",
    "efield_appendix": "Figure E.2",
}
APP_TAB_LABEL = {
    "param_inventory": "Table A.1",
    "stack_appendix": "Table B.1",
    "convergence_records": "Table C.1",
    "rejected": "Table D.1",
}


def title_page(doc):
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(C.TITLE, style="Title")
    p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(C.SUBTITLE)
    r.font.size = Pt(12)
    r.font.italic = True
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(30)

    tbl = doc.add_table(rows=0, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for field, value in C.TITLE_PAGE:
        row = tbl.add_row()
        no_split(row)
        for i, (txt, w, bold) in enumerate(
                ((field, 1.6, True), (value, 3.4, False))):
            c = row.cells[i]
            c.width = Inches(w)
            c.text = ""
            par = c.paragraphs[0]
            par.style = doc.styles["Front Matter Entry"]
            run = par.add_run(txt)
            run.bold = bold
            if txt.startswith("["):
                run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            clear_cell_borders(c)
        for i, w in enumerate((1.6, 3.4)):
            tbl.columns[i].width = Inches(w)

    doc.add_paragraph().paragraph_format.space_after = Pt(24)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Abstract")
    r.bold = True
    r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(C.ABSTRACT)
    r.italic = True
    r.font.color.rgb = GREY
    p.paragraph_format.space_after = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Cm(1.2)
    p.paragraph_format.right_indent = Cm(1.2)
    r = p.add_run(C.TITLE_NOTE)
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY

    doc.add_page_break()


def list_block(doc, heading, entries, page_map, break_before=False):
    """entries: list of (level, text, key)."""
    h = doc.add_heading(heading, level=1)
    h.paragraph_format.space_before = Pt(0)
    # Attaching the break to the heading avoids the stray empty page that an
    # explicit break paragraph produces when the previous block ends flush with
    # the bottom margin.
    h.paragraph_format.page_break_before = break_before
    for item in entries:
        if len(item) == 3:
            level, text, key = item
        else:
            level, text, key = 1, item[0], item[1]
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3.5)
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent = Cm(0.75 * (level - 1))
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(16.0 - 0.75 * (level - 1)), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(text)
        r.font.size = Pt(10.5 if level == 1 else 10)
        if level == 1:
            r.bold = True
        pn = page_map.get(key)
        p.add_run("\t")
        rr = p.add_run(str(pn) if pn else "")
        rr.font.size = Pt(10)


def main(out_path, page_map):
    doc = Document()
    build_styles(doc)
    page_setup(doc)

    # Section 1: front matter, roman-free but unnumbered header
    decorate_section(doc.sections[0], None, numbering=False)
    title_page(doc)

    # Build the body into a scratch pass first so that the contents can be
    # written before it. The Builder records headings, figures and tables.
    scratch = Document()
    build_styles(scratch)
    sb = Builder(scratch, page_map)
    sb.render(C.BODY, C.FIGURES, C.TABLES)
    for letter, title, blocks in C.APPENDICES:
        sb.appendix_heading(letter, title)
        sb.render(blocks, C.FIGURES, C.TABLES,
                  fig_label=lambda k: APP_FIG_LABEL.get(k),
                  tab_label=lambda k: APP_TAB_LABEL.get(k))

    toc_entries = list(sb.toc)
    toc_entries.insert(len(
        [t for t in sb.toc if not t[2].startswith("app:")]),
        (1, "References", "refs"))
    fig_entries = [(1, "%s  %s" % (l, s), k) for l, s, k in sb.figures]
    tab_entries = [(1, "%s  %s" % (l, s), k) for l, s, k in sb.tables]

    # ---- contents
    if STATIC_TOC:
        list_block(doc, "Contents", toc_entries, page_map)
    else:
        doc.add_heading("Contents", level=1)
        p = doc.add_paragraph()
        add_field(p, ' TOC \\o "1-2" \\h \\z \\u ')

    list_block(doc, "List of figures", fig_entries, page_map, break_before=True)
    list_block(doc, "List of tables", tab_entries, page_map)

    nom = doc.add_heading("Nomenclature", level=1)
    tbl = doc.add_table(rows=0, cols=2)
    tbl.autofit = False
    for term, meaning in C.NOMENCLATURE:
        row = tbl.add_row()
        no_split(row)
        for i, (txt, w, bold) in enumerate(((term, 1.35, True), (meaning, 4.85, False))):
            c = row.cells[i]
            c.width = Inches(w)
            c.text = ""
            par = c.paragraphs[0]
            par.style = doc.styles["Front Matter Entry"]
            par.add_run(txt).bold = bold
            clear_cell_borders(c)
        for i, w in enumerate((1.35, 4.85)):
            tbl.columns[i].width = Inches(w)

    # ---- main matter in a fresh section so page numbering can start here
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    page_setup(doc)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    decorate_section(sec, "CST reconstruction of a varactor-tunable 26 GHz reflectarray unit cell")
    sectPr = sec._sectPr
    pgNum = _el("w:pgNumType", **{"w:start": "1"})
    sectPr.append(pgNum)

    b = Builder(doc, page_map)
    b.render(C.BODY, C.FIGURES, C.TABLES)

    # ---- references
    doc.add_heading("References", level=1)
    for i, ref in enumerate(C.REFERENCES, 1):
        p = doc.add_paragraph(style="Reference Entry")
        p.add_run("[%d]\t" % i)
        p.add_run(ref)
    p = doc.add_paragraph(style="Small Note")
    p.add_run(C.REFERENCES_NOTE)

    # ---- appendices
    first_appendix = True
    for letter, title, blocks in C.APPENDICES:
        h = b.appendix_heading(letter, title)
        if first_appendix:
            h.paragraph_format.page_break_before = True
            first_appendix = False
        b.render(blocks, C.FIGURES, C.TABLES,
                 fig_label=lambda k: APP_FIG_LABEL.get(k),
                 tab_label=lambda k: APP_TAB_LABEL.get(k))

    doc.save(out_path)
    meta = dict(
        toc=[list(t) for t in toc_entries],
        figures=[list(f) for f in fig_entries],
        tables=[list(t) for t in tab_entries],
    )
    with open(os.path.splitext(out_path)[0] + ".outline.json", "w") as fh:
        json.dump(meta, fh, indent=1)
    print("wrote %s" % out_path)
    print("  headings %d, figures %d, tables %d"
          % (len(toc_entries), len(fig_entries), len(tab_entries)))


if __name__ == "__main__":
    out = sys.argv[1]
    pm = {}
    if len(sys.argv) > 2 and os.path.exists(sys.argv[2]):
        pm = json.load(open(sys.argv[2]))
    main(out, pm)
